"""
Ingestion pipeline — file → sections → chunks → embeddings → Qdrant.

This is the single code path behind POST /ingest and the one an evaluation
harness calls directly, so what gets measured is what gets served:

    extract_and_chunk()   pure CPU work, no models, no network — testable alone
    index_document()      extract_and_chunk + embed + BM25 + upsert (async)

Identity is content-derived. doc_id = uuid5(deal_id, sha256(file bytes)) and
every point id = uuid5(chunk_id), so re-uploading the same file into the same
deal overwrites its own points instead of duplicating them (the old ids were
uuid4 doc_ids and `hash(chunk_id)`, which Python salts per process).

Payload fields written per child chunk (see _child_payload):
    chunk_id, deal_id, doc_id, text, source_file, file_type, document_category,
    section_heading, page_number (None when the format has no pages: .txt,
    .docx, .xlsx), clause_id, is_table, content_type, table_id (tables only),
    table_representation (converter tables only), parent_chunk_id (prose only),
    chunk_index, token_count, is_current_version, contains_pii, risk_signals,
    supersedes_doc_id, superseded_by, is_redline, plus format extras
    (sheet_name, currency, scale_factor, scale_label, slide_number, page_range,
    metrics, redline_base_doc_id).
"""

from __future__ import annotations

import asyncio
import hashlib
import os
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Awaitable, Callable

import numpy as np

from src.data_processing.document_classifier import DocumentClassifier
from src.data_processing.excel_normalizer import ExcelNormalizer
from src.data_processing.financial_table_converter import (
    FinancialTableConverter,
    frame_from_rows,
    representation_content_type,
)
from src.data_processing.pii_detector import PIIDetector
from src.data_processing.risk_signal_extractor import RiskSignalExtractor
from src.data_processing.semantic_chunker import (
    DEFAULT_PARENT_TARGET_TOKENS,
    SemanticChunker,
)
from src.data_processing.structural_chunker import StructuralChunker
from src.data_processing.text_processor import parse_text_sections
from src.vector_db.constants import (
    COLLECTION_NAME,
    PARENT_COLLECTION_NAME,
    QDRANT_BATCH_SIZE,
)
from src.utils.logger import setup_logger

logger = setup_logger(__name__)

# .xls is deliberately absent: openpyxl cannot read the legacy binary format.
SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".pptx", ".xlsx", ".txt"})
_OFFICE_ZIP_EXTENSIONS = frozenset({".docx", ".pptx", ".xlsx"})

# Zip-bomb guards for Office files (which are zip archives). The limits are
# checked against the central directory before any parser opens the file;
# zipfile then refuses to inflate an entry beyond its declared size, so a
# forged directory cannot smuggle more data past this check.
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ENTRY_COMPRESSION_RATIO = 200
_RATIO_CHECK_MIN_BYTES = 10 * 1024 * 1024

# Fixed namespace for every ingestion id. Changing it changes every doc_id and
# point id, i.e. forces a full re-index.
INGEST_NAMESPACE = uuid.UUID("6f1c2a52-5d0e-4f8e-9a57-3c0b8e2d9a11")

_CONTENT_SAMPLE_CHARS = 2000

# Non-core metadata copied from a chunk into its payload when present.
_PASSTHROUGH_KEYS = (
    "sheet_name",
    "currency",
    "scale_factor",
    "scale_label",
    "slide_number",
    "page_range",
    "table_representation",
    "metrics",
    "redline_base_doc_id",
)

EmbedFn = Callable[[list[str]], Awaitable[Any]]
SparseFn = Callable[[list[str]], list]


# ==============================================================================
# Errors — messages are safe to return to an API client
# ==============================================================================


class IngestionError(Exception):
    """An ingestion failure whose message is safe to show a client."""

    status_code = 422


class UnsupportedFileTypeError(IngestionError):
    """The file extension is not one the pipeline can process."""

    status_code = 400


class UnsafeDocumentError(IngestionError):
    """The file exceeds a safety limit (e.g. a decompression bomb)."""

    status_code = 413


class DocumentExtractionError(IngestionError):
    """The file could not be parsed as the format its extension claims."""

    status_code = 422


class NoExtractableContentError(IngestionError):
    """Parsing succeeded but produced nothing to index."""

    status_code = 422


# ==============================================================================
# Result container
# ==============================================================================


@dataclass
class ExtractedDocument:
    """Everything extract_and_chunk() produces, before any embedding."""
    doc_id: str
    deal_id: str
    filename: str
    extension: str
    document_category: str
    content_sha256: str
    chunks: list[dict] = field(default_factory=list)    # child payloads
    parents: list[dict] = field(default_factory=list)   # parent payloads
    risk_signals: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    has_redline: bool = False


# ==============================================================================
# Identity and input validation
# ==============================================================================


def sanitize_filename(name: str | None) -> str:
    """
    Reduces a client-supplied filename to a safe display basename.

    The result is used only as metadata (source_file, classification); the
    upload is written to disk under a generated name, so this is not the only
    thing standing between a client and the filesystem.

    Args:
        name: Raw filename from the multipart upload.

    Returns:
        Basename with path components, control characters and Windows-reserved
        characters removed.

    Raises:
        UnsupportedFileTypeError: If nothing usable remains, or it is a dotfile.
    """
    if not name:
        raise UnsupportedFileTypeError("A filename is required")

    # Split on both separators: Path() only understands the host's own.
    base = re.split(r"[\\/]", name)[-1]
    base = "".join(ch for ch in base if ch.isprintable())
    base = re.sub(r'[<>:"|?*]', "_", base).strip()

    if not base or base in {".", ".."} or base.startswith("."):
        raise UnsupportedFileTypeError("Invalid filename")

    if len(base) > 200:
        stem, ext = os.path.splitext(base)
        base = stem[: 200 - len(ext)] + ext
    return base


def compute_content_hash(file_path: str) -> str:
    """
    SHA-256 of a file's bytes, streamed.

    Args:
        file_path: Path to the file.

    Returns:
        Hex digest.
    """
    digest = hashlib.sha256()
    with open(file_path, "rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def compute_doc_id(deal_id: str, content_sha256: str) -> str:
    """
    Deterministic document id: same bytes in the same deal → same id.

    Args:
        deal_id: Deal the document belongs to.
        content_sha256: SHA-256 hex digest of the file bytes.

    Returns:
        UUID string.
    """
    return str(uuid.uuid5(INGEST_NAMESPACE, f"doc:{deal_id}:{content_sha256}"))


def point_id_for(chunk_id: str) -> str:
    """
    Deterministic Qdrant point id for a chunk_id (stable across processes).

    Args:
        chunk_id: Chunk identifier.

    Returns:
        UUID string accepted by Qdrant as a point id.
    """
    return str(uuid.uuid5(INGEST_NAMESPACE, f"point:{chunk_id}"))


def _max_uncompressed_bytes() -> int:
    """Total uncompressed-size cap for Office archives (env MAX_UNCOMPRESSED_MB)."""
    return int(os.getenv("MAX_UNCOMPRESSED_MB", "200")) * 1024 * 1024


def validate_office_archive(file_path: str, extension: str) -> None:
    """
    Rejects .docx/.pptx/.xlsx files that would decompress to an unsafe size.

    Args:
        file_path: Path to the uploaded file.
        extension: Lower-case extension including the dot.

    Raises:
        DocumentExtractionError: If the file is not a zip archive at all.
        UnsafeDocumentError: If entry count, total size or ratio exceed limits.
    """
    if extension not in _OFFICE_ZIP_EXTENSIONS:
        return

    try:
        archive = zipfile.ZipFile(file_path)
    except zipfile.BadZipFile as e:
        raise DocumentExtractionError(
            f"The file is not a valid {extension} document"
        ) from e

    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_ARCHIVE_ENTRIES:
            raise UnsafeDocumentError("Document archive has too many entries")

        total = sum(info.file_size for info in infos)
        if total > _max_uncompressed_bytes():
            raise UnsafeDocumentError(
                "Document expands beyond the allowed uncompressed size"
            )

        for info in infos:
            if (
                info.file_size > _RATIO_CHECK_MIN_BYTES
                and info.compress_size > 0
                and info.file_size / info.compress_size > MAX_ENTRY_COMPRESSION_RATIO
            ):
                raise UnsafeDocumentError("Document has an abnormal compression ratio")


# ==============================================================================
# Extraction
# ==============================================================================


def _content_sample(file_path: str, extension: str) -> str:
    """
    First ~2000 characters of text, for content-based classification.

    Best effort: classification falls back to filename/extension heuristics,
    so any failure here returns "".

    Args:
        file_path: Path to the file.
        extension: Lower-case extension including the dot.

    Returns:
        Text sample (possibly empty).
    """
    try:
        if extension == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read(_CONTENT_SAMPLE_CHARS)
        if extension == ".pdf":
            import fitz
            with fitz.open(file_path) as pdf:
                text = ""
                for page in pdf:
                    text += page.get_text("text")
                    if len(text) >= _CONTENT_SAMPLE_CHARS:
                        break
                return text[:_CONTENT_SAMPLE_CHARS]
        if extension == ".docx":
            import docx
            document = docx.Document(file_path)
            return "\n".join(p.text for p in document.paragraphs[:80])[:_CONTENT_SAMPLE_CHARS]
        if extension == ".pptx":
            from pptx import Presentation
            parts = []
            for slide in list(Presentation(file_path).slides)[:5]:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            return "\n".join(parts)[:_CONTENT_SAMPLE_CHARS]
    except Exception as e:
        logger.warning("Content sample extraction failed", extra={"error": str(e)})
    return ""


def _table_sections(
    rows: list[list[Any]] | None,
    fallback_text: str,
    heading: str,
    page_number: int | None,
    table_id: str,
    extra: dict | None = None,
) -> list[dict]:
    """
    Turns one table into sections sharing a table_id.

    When the rows parse as a numeric table (header + labelled rows), the
    FinancialTableConverter's representations are emitted — narrative,
    row_by_row, markdown, and metrics_summary when a metric is computable.
    Otherwise the table is indexed once, verbatim, as content_type table_text.

    Args:
        rows: Header row followed by data rows, or None if unavailable.
        fallback_text: Verbatim table text for the non-numeric case.
        heading: Section heading the table sits under.
        page_number: Page (or slide) number, None when the format has none.
        table_id: Identifier shared by every representation.
        extra: Additional section keys (slide_number, page_range, ...).

    Returns:
        List of section dicts.
    """
    extra = dict(extra or {})
    base = {
        "section_heading": heading,
        "page_number": page_number,
        "section_type": "table",
        "is_table": 1,
        "table_id": table_id,
        **extra,
    }

    table_df = frame_from_rows(rows[0], rows[1:]) if rows and len(rows) >= 2 else None
    if table_df is None:
        if not fallback_text.strip():
            return []
        return [{**base, "text": fallback_text, "content_type": "table_text"}]

    meta = ExcelNormalizer().detect_scale([str(c) for c in rows[0]] + [heading])
    representations = FinancialTableConverter().generate_all_representations(
        df=table_df, meta=meta, table_id=table_id, source_metadata={},
    )

    sections = []
    prefix = f"{heading}\n" if heading else ""
    for rep in representations:
        rep_type = rep["table_representation"]
        if rep_type == "metrics_summary" and not rep.get("metrics"):
            continue  # nothing computable — the summary would be a bare title
        section = {
            **base,
            "text": prefix + rep["text"],
            "table_representation": rep_type,
            "content_type": representation_content_type(rep_type),
            "currency": meta.currency,
            "scale_factor": meta.scale_factor,
            "scale_label": meta.scale_label,
        }
        if rep.get("metrics"):
            section["metrics"] = rep["metrics"]
        sections.append(section)
    return sections


def extract_sections(
    file_path: str,
    extension: str,
    doc_id: str,
    deal_id: str,
    document_category: str,
) -> tuple[list[dict], list[str], bool]:
    """
    Runs the format-specific processor and normalises its output to sections.

    Args:
        file_path: Path to the file on disk.
        extension: Lower-case extension including the dot.
        doc_id: Document id (for processor logging and redline linkage).
        deal_id: Deal id (table_ids are "{deal_id}_{doc_id}_tNNN").
        document_category: Category (legal PDFs use clause segmentation).

    Returns:
        Tuple of (sections, warnings, has_redline).
    """
    sections: list[dict] = []
    warnings: list[str] = []
    has_redline = False
    table_seq = 0

    def next_table_id() -> str:
        nonlocal table_seq
        table_id = f"{deal_id}_{doc_id}_t{table_seq:03d}"
        table_seq += 1
        return table_id

    if extension == ".pdf":
        from src.data_processing.pdf_processor import PDFProcessor
        processor = PDFProcessor(legal_mode=document_category == "legal")
        for s in processor.process(file_path, doc_id):
            if s.is_table:
                sections.extend(_table_sections(
                    rows=s.table_rows,
                    fallback_text=s.text,
                    heading=s.section_heading,
                    page_number=s.page_number,
                    table_id=next_table_id(),
                    extra={"page_range": s.page_range} if s.page_range else None,
                ))
            else:
                sections.append({
                    "text": s.text,
                    "section_heading": s.section_heading,
                    "page_number": s.page_number,
                    "section_type": s.section_type,
                    "clause_id": s.clause_id,
                })

    elif extension == ".docx":
        from src.data_processing.docx_processor import process_docx_with_versions
        clean_chunks, redline_chunks = process_docx_with_versions(file_path, doc_id)
        for chunk in clean_chunks:
            if chunk.metadata.get("table_rows"):
                sections.extend(_table_sections(
                    rows=chunk.metadata["table_rows"],
                    fallback_text=chunk.text,
                    heading=chunk.section_heading,
                    page_number=None,
                    table_id=next_table_id(),
                ))
            else:
                sections.append({
                    "text": chunk.text,
                    "section_heading": chunk.section_heading,
                    "page_number": None,
                    "section_type": "text",
                })
        # Redlines are indexed alongside the clean text (is_redline=1) for the
        # paragraphs that actually carry tracked changes.
        for chunk in redline_chunks:
            has_redline = True
            sections.append({
                "text": chunk.text,
                "section_heading": chunk.section_heading,
                "page_number": None,
                "section_type": "text",
                "content_type": "redline",
                "is_redline": 1,
                "redline_base_doc_id": chunk.redline_base_doc_id,
            })

    elif extension == ".pptx":
        from src.data_processing.pptx_processor import PPTXProcessor
        processor = PPTXProcessor()
        for s in processor.to_text_chunks(processor.process(file_path, doc_id)):
            if s.get("is_table"):
                sections.extend(_table_sections(
                    rows=s.get("table_rows"),
                    fallback_text=s["text"],
                    heading=s.get("section_heading", ""),
                    page_number=s.get("page_number"),
                    table_id=next_table_id(),
                    extra={"slide_number": s.get("slide_number")},
                ))
            else:
                sections.append(s)

    elif extension == ".xlsx":
        from src.data_processing.excel_processor import ExcelProcessor
        processor = ExcelProcessor()
        sheets = processor.process(file_path, doc_id, table_id_prefix=f"{deal_id}_{doc_id}")
        sections = processor.to_chunks(sheets)
        warnings.extend(
            f"Sheet '{f['sheet_name']}' could not be processed"
            for f in processor.failed_sheets
        )

    elif extension == ".txt":
        with open(file_path, "r", encoding="utf-8-sig", errors="replace") as f:
            text = f.read()
        for s in parse_text_sections(text):
            if s["is_table"]:
                sections.extend(_table_sections(
                    rows=None,  # whitespace-aligned text tables stay verbatim
                    fallback_text=s["text"],
                    heading=s["section_heading"],
                    page_number=None,
                    table_id=next_table_id(),
                ))
            else:
                sections.append(s)

    else:
        raise UnsupportedFileTypeError(f"Unsupported file type: {extension}")

    return sections, warnings, has_redline


def _child_payload(
    chunk,
    chunk_id: str,
    index: int,
    doc: ExtractedDocument,
    is_current_version: bool,
    supersedes_doc_id: str | None,
    contains_pii: int,
    risk_signals: list[str],
    parent_chunk_id: str | None,
) -> dict:
    """
    Builds the Qdrant payload for one child chunk.

    Args:
        chunk: SemanticChunk.
        chunk_id: Stable chunk id.
        index: Position of the chunk in the document.
        doc: Document being built (for ids, filename, category).
        is_current_version: Version flag.
        supersedes_doc_id: Document this one replaces, if any.
        contains_pii: 0/1 PII flag.
        risk_signals: Risk signal types detected in this chunk.
        parent_chunk_id: Parent chunk id, or None for tables/redlines.

    Returns:
        Payload dict.
    """
    meta = chunk.metadata
    payload = {
        "chunk_id": chunk_id,
        "deal_id": doc.deal_id,
        "doc_id": doc.doc_id,
        "text": chunk.text,
        "source_file": doc.filename,
        "file_type": doc.extension.lstrip("."),
        "document_category": doc.document_category,
        "section_heading": chunk.section_heading or "",
        "page_number": chunk.page_number,
        "clause_id": chunk.clause_id,
        "is_table": int(meta.get("is_table", 0)),
        "content_type": meta.get("content_type", "text"),
        "chunk_index": index,
        "token_count": chunk.token_count,
        "is_current_version": 1 if is_current_version else 0,
        "contains_pii": contains_pii,
        "risk_signals": risk_signals,
        "supersedes_doc_id": supersedes_doc_id or "",
        "superseded_by": "",  # stamped later if a newer version replaces this doc
        "is_redline": int(meta.get("is_redline", 0)),
    }
    if meta.get("table_id"):
        payload["table_id"] = meta["table_id"]
    if parent_chunk_id:
        payload["parent_chunk_id"] = parent_chunk_id
    for key in _PASSTHROUGH_KEYS:
        if meta.get(key) is not None:
            payload[key] = meta[key]
    return payload


def extract_and_chunk(
    file_path: str,
    filename: str,
    deal_id: str,
    document_category: str | None = None,
    *,
    is_current_version: bool = True,
    supersedes_doc_id: str | None = None,
    doc_id: str | None = None,
) -> ExtractedDocument:
    """
    Extracts, chunks and annotates a document — everything short of embedding.

    Synchronous and CPU-bound (parsing, tokenising, PII and risk regexes):
    call it through asyncio.to_thread from async code. No models are loaded
    and no network is used, so it is directly unit-testable.

    Args:
        file_path: Path to the file on disk (any name; the format is taken
                   from `filename`).
        filename: Original (sanitised) filename — stored as source_file.
        deal_id: Deal the document belongs to.
        document_category: Category override; auto-classified when None.
        is_current_version: Version flag written to every payload.
        supersedes_doc_id: Document this one replaces, if any.
        doc_id: Explicit doc id; by default derived from deal_id + content hash.

    Returns:
        ExtractedDocument with child and parent payloads (no vectors).

    Raises:
        UnsupportedFileTypeError: Unknown extension.
        UnsafeDocumentError: Office archive over the safety limits.
        DocumentExtractionError: The file could not be parsed.
        NoExtractableContentError: Parsing produced nothing to index.
    """
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type: {extension or '(none)'}. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    validate_office_archive(file_path, extension)

    content_sha256 = compute_content_hash(file_path)
    doc_id = doc_id or compute_doc_id(deal_id, content_sha256)

    if document_category is None:
        document_category = DocumentClassifier().classify(
            file_name=filename,
            file_type=extension.lstrip("."),
            content_sample=_content_sample(file_path, extension),
        )

    try:
        sections, warnings, has_redline = extract_sections(
            file_path, extension, doc_id, deal_id, document_category
        )
    except IngestionError:
        raise
    except Exception as e:
        logger.error(
            "Document extraction failed",
            extra={"doc_id": doc_id, "file_name": filename, "error": str(e)},
            exc_info=True,
        )
        raise DocumentExtractionError(
            f"The {extension} file could not be parsed"
        ) from e

    for warning in warnings:
        logger.error("Partial extraction failure", extra={"doc_id": doc_id, "warning": warning})

    doc = ExtractedDocument(
        doc_id=doc_id,
        deal_id=deal_id,
        filename=filename,
        extension=extension,
        document_category=document_category,
        content_sha256=content_sha256,
        warnings=warnings,
        has_redline=has_redline,
    )

    if not sections:
        raise NoExtractableContentError("No text could be extracted from the document")

    structural_chunks = StructuralChunker().chunk(sections)
    children, parents = SemanticChunker().chunk_with_parents(
        structural_chunks, parent_target_tokens=DEFAULT_PARENT_TARGET_TOKENS
    )
    if not children:
        raise NoExtractableContentError("No text could be extracted from the document")

    # PII and risk signals are regex scans — no LLM, no added latency. Signals
    # are written per chunk AND aggregated per document for the risk dashboard.
    pii_detector = PIIDetector()
    risk_extractor = RiskSignalExtractor()
    aggregated_risk: dict[str, dict] = {}
    parent_ids = [f"{deal_id}_{doc_id}_p{j:04d}" for j in range(len(parents))]
    parent_pii = [0] * len(parents)
    parent_children: list[list[str]] = [[] for _ in parents]

    for i, chunk in enumerate(children):
        chunk_id = f"{deal_id}_{doc_id}_{i:04d}"
        contains_pii = pii_detector.detect(chunk.text).contains_pii

        risk = risk_extractor.extract(
            chunk.text, file_name=filename, document_category=document_category
        )
        for detail in risk.signal_details:
            entry = aggregated_risk.setdefault(
                detail["signal_type"],
                {
                    "signal_type": detail["signal_type"],
                    "match_count": 0,
                    "sample_matches": [],
                    "page_number": chunk.page_number,
                },
            )
            entry["match_count"] += detail["match_count"]
            for sample in detail["sample_matches"]:
                if sample and len(entry["sample_matches"]) < 3:
                    entry["sample_matches"].append(sample)

        parent_index = chunk.metadata.get("parent_index")
        parent_chunk_id = None
        if parent_index is not None:
            parent_chunk_id = parent_ids[parent_index]
            parent_pii[parent_index] = max(parent_pii[parent_index], contains_pii)
            parent_children[parent_index].append(chunk_id)

        doc.chunks.append(_child_payload(
            chunk, chunk_id, i, doc,
            is_current_version=is_current_version,
            supersedes_doc_id=supersedes_doc_id,
            contains_pii=contains_pii,
            risk_signals=risk.signals,
            parent_chunk_id=parent_chunk_id,
        ))

    for j, parent in enumerate(parents):
        doc.parents.append({
            "chunk_id": parent_ids[j],
            "deal_id": deal_id,
            "doc_id": doc_id,
            "text": parent.text,
            "source_file": filename,
            "document_category": document_category,
            "section_heading": parent.section_heading,
            "page_number": parent.page_number,
            "token_count": parent.token_count,
            "child_chunk_ids": parent_children[j],
            "is_current_version": 1 if is_current_version else 0,
            # A parent is shown with any of its children, so it is PII if any
            # child is — otherwise parent expansion would leak what the child
            # filter withheld.
            "contains_pii": parent_pii[j],
        })

    doc.risk_signals = list(aggregated_risk.values())

    logger.info(
        "Document extracted and chunked",
        extra={
            "doc_id": doc_id,
            "file_name": filename,
            "category": document_category,
            "sections": len(sections),
            "chunks": len(doc.chunks),
            "parents": len(doc.parents),
            "tables": len({c["table_id"] for c in doc.chunks if c.get("table_id")}),
        },
    )
    return doc


# ==============================================================================
# Indexing
# ==============================================================================


def _doc_filter(deal_id: str, doc_id: str, keep_ids: list[str] | None = None):
    """Filter for one document's points, optionally excluding the given ids."""
    from qdrant_client.models import FieldCondition, Filter, HasIdCondition, MatchValue

    return Filter(
        must=[
            FieldCondition(key="deal_id", match=MatchValue(value=deal_id)),
            FieldCondition(key="doc_id", match=MatchValue(value=doc_id)),
        ],
        must_not=[HasIdCondition(has_id=keep_ids)] if keep_ids else None,
    )


async def _upsert_batches(client, collection_name: str, points: list) -> None:
    """Upserts points in QDRANT_BATCH_SIZE batches, waiting for each."""
    for start in range(0, len(points), QDRANT_BATCH_SIZE):
        await client.upsert(
            collection_name=collection_name,
            points=points[start:start + QDRANT_BATCH_SIZE],
            wait=True,
        )


async def index_document(
    file_path: str,
    filename: str,
    deal_id: str,
    document_category: str | None = None,
    *,
    is_current_version: bool = True,
    supersedes_doc_id: str | None = None,
    client=None,
    collection_name: str = COLLECTION_NAME,
    parent_collection_name: str = PARENT_COLLECTION_NAME,
    embed_fn: EmbedFn | None = None,
    sparse_fn: SparseFn | None = None,
    ensure_collections: bool = False,
) -> dict:
    """
    Extracts, chunks, embeds and upserts one document. Idempotent per content.

    Steps: extract_and_chunk (worker thread) → dense embeddings → batched BM25
    (worker thread) → upsert children and parents → delete this document's
    stale points (ids no longer produced, e.g. after a chunking change).
    Re-indexing identical bytes therefore replaces rather than duplicates.

    If any upsert fails, every point for this (deal_id, doc_id) is deleted
    from both collections before the error propagates, so a document is
    either fully indexed or absent — never half-indexed.

    Args:
        file_path: Path to the file on disk.
        filename: Original filename (stored as source_file; sets the format).
        deal_id: Deal the document belongs to.
        document_category: Category override; auto-classified when None.
        is_current_version: Version flag for every point.
        supersedes_doc_id: Document this one replaces (the caller retires it
                           with mark_superseded()).
        client: AsyncQdrantClient; defaults to the application singleton.
        collection_name: Child collection.
        parent_collection_name: Parent collection.
        embed_fn: async (texts) -> array (n, dim); defaults to bge-m3.
        sparse_fn: (texts) -> list[SparseVector]; defaults to batched BM25.
        ensure_collections: Create collections/indexes first if missing.

    Returns:
        Dict with doc_id, deal_id, filename, document_category, chunks_created,
        parent_chunks_created, table_count, risk_signals, has_redline, warnings,
        content_sha256 and previously_indexed.

    Raises:
        IngestionError subclasses for client-attributable failures; other
        exceptions (Qdrant, embedding) propagate after rollback.
    """
    from qdrant_client.models import PointStruct

    if client is None:
        from src.vector_db.qdrant_client import get_qdrant_client
        client = get_qdrant_client()
    if embed_fn is None:
        from src.vector_db.reranker import embed_texts_async
        embed_fn = embed_texts_async
    if sparse_fn is None:
        from src.vector_db.hybrid_search import compute_sparse_bm25_batch
        sparse_fn = compute_sparse_bm25_batch
    if ensure_collections:
        from src.vector_db.collection_manager import setup_collections
        await setup_collections(client, collection_name, parent_collection_name)

    # Parsing, chunking, PII and risk scans are CPU-bound; running them on the
    # event loop froze every other request on the single-worker API.
    doc = await asyncio.to_thread(
        extract_and_chunk,
        file_path,
        filename,
        deal_id,
        document_category,
        is_current_version=is_current_version,
        supersedes_doc_id=supersedes_doc_id,
    )

    texts = [c["text"] for c in doc.chunks]
    dense = np.asarray(await embed_fn(texts), dtype=np.float32)
    sparse = await asyncio.to_thread(sparse_fn, texts)

    child_points = [
        PointStruct(
            id=point_id_for(payload["chunk_id"]),
            vector={"dense": dense[i].tolist(), "sparse": sparse[i]},
            payload=payload,
        )
        for i, payload in enumerate(doc.chunks)
    ]

    # The parent collection requires a dense vector, but parents are fetched by
    # id, never searched — the normalised mean of their children's vectors is a
    # meaningful placeholder that costs no extra embedding pass.
    child_index = {c["chunk_id"]: i for i, c in enumerate(doc.chunks)}
    parent_points = []
    for parent in doc.parents:
        rows = [child_index[cid] for cid in parent["child_chunk_ids"]]
        vector = dense[rows].mean(axis=0)
        norm = float(np.linalg.norm(vector)) or 1.0
        parent_points.append(PointStruct(
            id=point_id_for(parent["chunk_id"]),
            vector={"dense": (vector / norm).tolist()},
            payload=parent,
        ))

    existing = await client.count(
        collection_name=collection_name,
        count_filter=_doc_filter(deal_id, doc.doc_id),
        exact=True,
    )

    try:
        await _upsert_batches(client, collection_name, child_points)
        await _upsert_batches(client, parent_collection_name, parent_points)
    except Exception:
        logger.error(
            "Upsert failed; rolling back document",
            extra={"doc_id": doc.doc_id, "deal_id": deal_id},
            exc_info=True,
        )
        for name in (collection_name, parent_collection_name):
            try:
                await client.delete(
                    collection_name=name,
                    points_selector=_doc_filter(deal_id, doc.doc_id),
                    wait=True,
                )
            except Exception as rollback_error:
                logger.error(
                    "Rollback delete failed",
                    extra={"doc_id": doc.doc_id, "collection": name,
                           "error": str(rollback_error)},
                )
        raise

    # Remove points this document no longer produces (only relevant on re-index).
    if existing.count:
        for name, points in (
            (collection_name, child_points),
            (parent_collection_name, parent_points),
        ):
            await client.delete(
                collection_name=name,
                points_selector=_doc_filter(deal_id, doc.doc_id, [p.id for p in points]),
                wait=True,
            )

    result = {
        "doc_id": doc.doc_id,
        "deal_id": deal_id,
        "filename": filename,
        "document_category": doc.document_category,
        "chunks_created": len(child_points),
        "parent_chunks_created": len(parent_points),
        "table_count": len({c["table_id"] for c in doc.chunks if c.get("table_id")}),
        "risk_signals": doc.risk_signals,
        "has_redline": doc.has_redline,
        "warnings": doc.warnings,
        "content_sha256": doc.content_sha256,
        "previously_indexed": bool(existing.count),
    }
    logger.info(
        "Document indexed",
        extra={k: v for k, v in result.items() if k not in ("risk_signals", "warnings")},
    )
    return result


async def mark_superseded(
    deal_id: str,
    superseded_doc_id: str,
    superseded_by: str,
    client=None,
    collection_name: str = COLLECTION_NAME,
    parent_collection_name: str = PARENT_COLLECTION_NAME,
) -> None:
    """
    Flips every chunk (and parent) of a superseded document to is_current_version=0.

    Without this, uploading a replacement document leaves both versions marked
    current, and retrieval's is_current_version=1 filter happily returns stale
    terms alongside the ones that replaced them.

    Failures are logged, not raised: the new document is already indexed and
    usable, and losing the retirement stamp degrades results rather than
    invalidating the upload.

    Args:
        deal_id: Deal scope, so one deal can never retire another deal's docs.
        superseded_doc_id: Document being retired.
        superseded_by: Document ID replacing it.
        client: AsyncQdrantClient; defaults to the application singleton.
        collection_name: Child collection.
        parent_collection_name: Parent collection.
    """
    if superseded_doc_id == superseded_by:
        # Re-uploading identical bytes yields the same doc_id; retiring it
        # would retire the document that was just indexed.
        logger.warning(
            "Ignoring self-supersession", extra={"doc_id": superseded_doc_id}
        )
        return

    if client is None:
        from src.vector_db.qdrant_client import get_qdrant_client
        client = get_qdrant_client()

    try:
        for name in (collection_name, parent_collection_name):
            await client.set_payload(
                collection_name=name,
                payload={"is_current_version": 0, "superseded_by": superseded_by},
                points=_doc_filter(deal_id, superseded_doc_id),
            )
        logger.info(
            "Superseded document retired",
            extra={
                "deal_id": deal_id,
                "superseded_doc_id": superseded_doc_id,
                "superseded_by": superseded_by,
            },
        )
    except Exception as e:
        logger.error(
            "Failed to retire superseded document",
            extra={
                "deal_id": deal_id,
                "superseded_doc_id": superseded_doc_id,
                "error": str(e),
            },
        )
