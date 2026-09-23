"""
DOCX processor — parallel clean + redline processing.

In M&A, negotiation history (redlines) is often the most diligence-relevant
content — struck indemnification caps, added representations, renegotiated
earnouts. Silently collapsing to final version discards critical findings.

Produces two parallel chunk sets:
- Clean: tracked changes accepted, is_redline=0
- Redline: additions (+) and deletions (~~strikethrough~~) inline, is_redline=1

Clean text is read from the XML rather than python-docx's `paragraph.text`:
that property only sees runs that are direct children of the paragraph, so a
tracked insertion (runs wrapped in <w:ins>) silently vanished from the
"accepted" text — exactly the negotiated language a reviewer needs.
"""

from pathlib import Path
from dataclasses import dataclass
from lxml import etree

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.utils.logger import setup_logger

logger = setup_logger(__name__)

# Word XML namespaces
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NSMAP = {"w": WORD_NS}

_W = f"{{{WORD_NS}}}"
# Content inside these elements is not part of the accepted document.
_REJECTED_CONTAINERS = {f"{_W}del", f"{_W}moveFrom"}


@dataclass
class Chunk:
    """A text chunk from a DOCX document."""
    text: str
    section_heading: str = ""
    page_number: int | None = None  # DOCX has no stable page numbers without rendering
    is_redline: int = 0
    redline_base_doc_id: str = ""
    content_type: str = "text"
    metadata: dict = None

    def __post_init__(self):
        if self.metadata is None:
            self.metadata = {}


def _accepted_text(element) -> str:
    """
    Text of a paragraph (or cell) with all tracked changes accepted.

    Insertions are included, deletions and move-sources are excluded.

    Args:
        element: lxml element for a <w:p> (or any container of paragraphs).

    Returns:
        Accepted text.
    """
    parts: list[str] = []
    for node in element.iter():
        tag = node.tag
        if tag not in (f"{_W}t", f"{_W}tab", f"{_W}br", f"{_W}cr"):
            continue
        rejected = False
        parent = node.getparent()
        while parent is not None and parent is not element:
            if parent.tag in _REJECTED_CONTAINERS:
                rejected = True
                break
            parent = parent.getparent()
        if rejected:
            continue
        if tag == f"{_W}t":
            parts.append(node.text or "")
        elif tag == f"{_W}tab":
            parts.append("\t")
        else:
            parts.append("\n")
    return "".join(parts)


def _is_heading(para: Paragraph) -> bool:
    """True for Heading N / Title styled paragraphs."""
    name = para.style.name if para.style is not None and para.style.name else ""
    return name.startswith("Heading") or name == "Title"


def _iter_block_items(doc: docx.Document):
    """
    Yields paragraphs and tables in document order.

    `doc.paragraphs` and `doc.tables` are separate lists, so iterating them
    independently loses where each table sits relative to its heading.

    Args:
        doc: Opened python-docx Document.

    Yields:
        Paragraph or Table objects.
    """
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == f"{_W}p":
            yield Paragraph(child, doc)
        elif child.tag == f"{_W}tbl":
            yield Table(child, doc)


def _table_rows(table: Table) -> list[list[str]]:
    """
    Accepted text of every cell, row by row.

    Args:
        table: python-docx Table.

    Returns:
        Rows of cell strings (merged cells repeat, as python-docx reports them).
    """
    rows = []
    for row in table.rows:
        cells = [
            " ".join(" ".join(_accepted_text(p) for p in cell._tc.iter(f"{_W}p")).split())
            for cell in row.cells
        ]
        if any(cells):
            rows.append(cells)
    return rows


def _extract_clean_text(doc: docx.Document) -> list[Chunk]:
    """
    Extracts clean text with tracked changes accepted.
    Insertions are included, deletions are excluded. Tables are emitted in
    place as table chunks carrying their rows in metadata["table_rows"].

    Args:
        doc: Opened python-docx Document.

    Returns:
        List of Chunks with clean text and section headings.
    """
    chunks = []
    current_heading = ""

    for block in _iter_block_items(doc):
        if isinstance(block, Table):
            rows = _table_rows(block)
            if not rows:
                continue
            chunks.append(Chunk(
                text="\n".join(" | ".join(r) for r in rows),
                section_heading=current_heading,
                is_redline=0,
                content_type="table_text",
                metadata={"is_table": 1, "table_rows": rows},
            ))
            continue

        text = _accepted_text(block._element).strip()
        if _is_heading(block):
            if text:
                current_heading = text
            continue

        if not text:
            continue

        chunks.append(Chunk(
            text=text,
            section_heading=current_heading,
            is_redline=0,
            content_type="text",
        ))

    return chunks


def _extract_redline_text(doc: docx.Document) -> list[Chunk]:
    """
    Extracts text with tracked changes shown inline:
    - Insertions marked with (+added text)
    - Deletions marked with (~~deleted text~~)

    Only paragraphs that actually contain tracked changes are returned —
    an unchanged paragraph is already covered by the clean set, and labelling
    it a redline would be false.

    Args:
        doc: Opened python-docx Document.

    Returns:
        List of Chunks with redline markup (is_redline=1).
    """
    chunks = []
    current_heading = ""

    for para in doc.paragraphs:
        # Detect headings
        if _is_heading(para):
            current_heading = _accepted_text(para._element).strip()
            continue

        # Parse the paragraph XML to find tracked changes
        para_xml = para._element
        parts = []
        has_changes = False

        for child in para_xml:
            tag = etree.QName(child.tag).localname if isinstance(child.tag, str) else ""

            if tag == "r":
                # Normal run
                text_els = child.findall(f"{{{WORD_NS}}}t")
                text = "".join(t.text or "" for t in text_els)
                if text:
                    parts.append(text)

            elif tag == "ins":
                # Insertion
                has_changes = True
                runs = child.findall(f".//{{{WORD_NS}}}t")
                text = "".join(t.text or "" for t in runs)
                if text:
                    parts.append(f"(+{text})")

            elif tag == "del":
                # Deletion
                has_changes = True
                del_runs = child.findall(f".//{{{WORD_NS}}}delText")
                text = "".join(t.text or "" for t in del_runs)
                if text:
                    parts.append(f"(~~{text}~~)")

        full_text = "".join(parts).strip()
        if not full_text or not has_changes:
            continue

        chunks.append(Chunk(
            text=full_text,
            section_heading=current_heading,
            is_redline=1,
            content_type="redline",
        ))

    return chunks


def process_docx_with_versions(
    path: str,
    doc_id: str,
) -> tuple[list[Chunk], list[Chunk]]:
    """
    Process DOCX into two parallel chunk sets.

    Returns:
        Tuple of (clean_chunks, redline_chunks).
        clean_chunks: Tracked changes accepted. is_redline=0.
        redline_chunks: Shows additions (+) and deletions (~~strikethrough~~) inline,
                        only for paragraphs with tracked changes.
                        is_redline=1, redline_base_doc_id=doc_id of clean version.

    Args:
        path: Absolute path to the DOCX file.
        doc_id: Document identifier for metadata.

    Raises:
        FileNotFoundError: If path does not exist.
        docx.opc.exceptions.PackageNotFoundError: If file is not a valid DOCX.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"DOCX not found: {path}")

    logger.info(
        "Processing DOCX with versions",
        extra={"path": path, "doc_id": doc_id},
    )

    doc = docx.Document(path)

    # Clean version — tracked changes accepted
    clean_chunks = _extract_clean_text(doc)

    # Redline version — shows tracked changes inline
    redline_chunks = _extract_redline_text(doc)

    # Set redline_base_doc_id on all redline chunks
    for chunk in redline_chunks:
        chunk.redline_base_doc_id = doc_id

    logger.info(
        "DOCX processing complete",
        extra={
            "doc_id": doc_id,
            "clean_chunks": len(clean_chunks),
            "redline_chunks": len(redline_chunks),
        },
    )

    return clean_chunks, redline_chunks
